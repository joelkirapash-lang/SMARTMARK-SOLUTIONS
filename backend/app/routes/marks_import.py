import csv
import io
import re

from flask import Blueprint, g, jsonify, request, send_file
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models import (Mark, Exam, Subject, User, Grade, Stream, GradeSubject, TeacherAssignment,
                        SCHOOL_ADMIN, TEACHER, EDUCATOR, STUDENT)
from app.routes.marks import staff_scope
from app.utils.audit import log_action
from app.utils.auth import require, scoped

bp = Blueprint("marks_import", __name__, url_prefix="/api")
ADMIN = SCHOOL_ADMIN
STAFF_VIEW = (SCHOOL_ADMIN, TEACHER, EDUCATOR)
MAX_ROWS = 2000
ADM_KEYS = {"admissionnumber", "admissionno", "admno", "admission", "adm", "admnumber", "admnum", "regno",
            "registrationnumber", "admissionnum"}
NAME_KEYS = {"name", "names", "studentname", "learnername", "fullname", "student", "learner", "pupil", "pupilname"}


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _read_table(f):
    name = (f.filename or "").lower()
    data = f.read()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [[_cell(c) for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported. In Excel choose Save As and pick .xlsx or .csv.")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            delim = csv.Sniffer().sniff(text[:2048], delimiters=",;\t").delimiter
        except csv.Error:
            delim = ","
        rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return [r for r in rows if any(c for c in r)]


def _split_header(h):
    """'Mathematics (50)' or 'Maths /50' -> ('Mathematics', 50.0)."""
    m = re.match(r"^(.*?)\s*[\(\[/]\s*(\d+(?:\.\d+)?)\s*[\)\]]?$", h.strip())
    if m and m.group(1).strip():
        return m.group(1).strip(), float(m.group(2))
    return h.strip(), None


def _context(src, need_exam):
    grade = scoped(Grade).filter_by(id=src.get("grade_id")).first()
    if grade is None:
        return None, (jsonify(error="Choose a grade."), 400)
    exam = None
    if need_exam:
        exam = scoped(Exam).filter_by(id=src.get("exam_id")).first()
        if exam is None:
            return None, (jsonify(error="Choose an exam."), 400)
    stream = None
    if src.get("stream_id"):
        stream = scoped(Stream).filter_by(id=src["stream_id"], grade_id=grade.id).first()
        if stream is None:
            return None, (jsonify(error="That stream does not belong to the selected grade."), 400)
    if g.role != ADMIN and not TeacherAssignment.query.filter_by(
            school_id=g.school_id, teacher_id=g.current_user.id, grade_id=grade.id).first():
        return None, (jsonify(error="You are not assigned to any subject in this grade."), 403)
    return (exam, grade, stream), None


@bp.post("/marks/import")
@require(*STAFF_VIEW)
def import_marks():
    ctx, err = _context(request.form, need_exam=True)
    if err:
        return err
    exam, grade, stream = ctx
    if exam.is_published and g.role != ADMIN:
        return jsonify(error="This exam is published, so marks are locked. Ask your administrator to make changes."), 403
    dry = request.form.get("dry_run", "1") != "0"
    try:
        default_max = float(request.form.get("max_score") or 100)
        if default_max <= 0:
            raise ValueError
    except ValueError:
        return jsonify(error="The out-of mark must be a number above 0."), 400
    f = request.files.get("file")
    if f is None:
        return jsonify(error="Choose a CSV or Excel file."), 400
    try:
        table = _read_table(f)
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except Exception:
        return jsonify(error="That file could not be read. Save it as .xlsx or .csv and try again."), 400
    if len(table) < 2:
        return jsonify(error="The file needs a header row and at least one learner."), 400
    if len(table) - 1 > MAX_ROWS:
        return jsonify(error=f"Please import {MAX_ROWS} rows or fewer at a time."), 400

    header = table[0]
    adm_i = name_i = None
    subj_cols = []
    for i, h in enumerate(header):
        n = _norm(h)
        if not n:
            continue
        if n in ADM_KEYS and adm_i is None:
            adm_i = i
        elif n in NAME_KEYS and name_i is None:
            name_i = i
        else:
            subj_cols.append((i, h))
    if adm_i is None and name_i is None:
        return jsonify(error="Add an 'Admission Number' column (or a 'Name' column) so each row can be matched to a learner."), 400

    by_key = {}
    for s in scoped(Subject).all():
        by_key[_norm(s.name)] = s
        by_key[_norm(s.code)] = s

    cols, used = [], set()
    for i, h in subj_cols:
        label, mx = _split_header(h)
        s = by_key.get(_norm(label))
        c = {"index": i, "header": h, "subject": s, "max": mx or default_max, "scope": None, "status": "ok", "reason": None}
        if s is None:
            c.update(status="unknown", reason="No subject with this name or code exists in your school.")
        elif s.id in used:
            c.update(status="blocked", reason="This subject appears twice. Only the first column is used.")
        else:
            c["scope"] = True if g.role == ADMIN else staff_scope(g.current_user, grade.id, s.id)
            if c["scope"] is None:
                c.update(status="blocked", reason=f"You are not assigned to {s.name} in {grade.name}.")
            else:
                used.add(s.id)
        cols.append(c)
    ok_cols = [c for c in cols if c["status"] == "ok"]

    q = scoped(User).filter_by(role=STUDENT, grade_id=grade.id, is_archived=False)
    if stream is not None:
        q = q.filter_by(stream_id=stream.id)
    roster = q.all()
    streams = {s.id: s.name for s in scoped(Stream).all()}
    adm_map = {_norm(s.admission_number): s for s in roster if s.admission_number}
    name_map = {}
    for s in roster:
        for key in {_norm(s.full_name()), _norm(f"{s.first_name} {s.second_name or ''}")}:
            name_map.setdefault(key, [])
            if s not in name_map[key]:
                name_map[key].append(s)

    rows_out, save_map = [], {}
    for ridx, row in enumerate(table[1:], start=2):
        get = lambda i: row[i] if i is not None and i < len(row) else ""
        adm, nm = get(adm_i), get(name_i)
        learner, problems = None, []
        if adm:
            learner = adm_map.get(_norm(adm))
            if learner is None:
                problems.append(f"Admission number {adm} is not in this class.")
        elif nm:
            hits = name_map.get(_norm(nm), [])
            if len(hits) == 1:
                learner = hits[0]
            elif len(hits) > 1:
                problems.append("More than one learner has this name. Add the admission number.")
            else:
                problems.append(f"No learner named {nm} in this class.")
        else:
            problems.append("This row has no admission number or name.")
        scores = {}
        if learner:
            for c in ok_cols:
                raw = get(c["index"])
                if raw == "":
                    continue
                name = c["subject"].name
                if c["scope"] is not True and learner.stream_id not in c["scope"]:
                    problems.append(f"{name}: you are not assigned to {streams.get(learner.stream_id, 'this stream')}.")
                    continue
                try:
                    v = float(raw.replace(",", "."))
                except ValueError:
                    problems.append(f"{name}: '{raw}' is not a number.")
                    continue
                if v < 0 or v > c["max"]:
                    problems.append(f"{name}: {raw} is outside 0 to {c['max']:g}.")
                    continue
                scores[c["subject"].id] = v
                save_map[(learner.id, c["subject"].id)] = (learner, c["subject"], v, c["max"])
        rows_out.append({"row": ridx, "learner": learner.full_name() if learner else (nm or None),
                         "admission_number": learner.admission_number if learner else (adm or None),
                         "matched": learner is not None, "scores": scores, "problems": problems})

    saved = 0
    if not dry and save_map:
        existing = {(m.student_id, m.subject_id): m for m in scoped(Mark).filter_by(exam_id=exam.id).all()}
        for (sid, subid), (learner, subj, v, mx) in save_map.items():
            m = existing.get((sid, subid))
            if m is None:
                m = Mark(school_id=g.school_id, exam_id=exam.id, student_id=sid, subject_id=subid, grade_id=grade.id)
                db.session.add(m)
            m.stream_id, m.score, m.max_score, m.entered_by_user_id = learner.stream_id, v, mx, g.current_user.id
            saved += 1
        log_action(g.school_id, g.current_user.id, "MARKS_IMPORTED",
                   f"{saved} marks · {grade.name}{(' ' + stream.name) if stream else ''} · {exam.name}")
        db.session.commit()

    return jsonify(
        dry_run=dry, saved=saved,
        subjects=[{"header": c["header"], "subject_id": c["subject"].id if c["subject"] else None,
                   "name": c["subject"].name if c["subject"] else None, "status": c["status"], "reason": c["reason"],
                   "max_score": c["max"]} for c in cols],
        rows=rows_out,
        summary={"rows": len(rows_out), "matched": sum(1 for r in rows_out if r["matched"]),
                 "marks": len(save_map), "problem_rows": sum(1 for r in rows_out if r["problems"])})


@bp.get("/marks/import-template")
@require(*STAFF_VIEW)
def import_template():
    ctx, err = _context(request.args, need_exam=False)
    if err:
        return err
    _, grade, stream = ctx
    if g.role == ADMIN:
        mapped = [r.subject_id for r in GradeSubject.query.filter_by(school_id=g.school_id, grade_id=grade.id)]
        subs = scoped(Subject).filter(Subject.id.in_(mapped)).all() if mapped else scoped(Subject).all()
        allowed_streams = True
    else:
        rows = TeacherAssignment.query.filter_by(school_id=g.school_id, teacher_id=g.current_user.id, grade_id=grade.id).all()
        subs = scoped(Subject).filter(Subject.id.in_({r.subject_id for r in rows})).all()
        allowed_streams = True if any(r.stream_id is None for r in rows) else {r.stream_id for r in rows}
    q = scoped(User).filter_by(role=STUDENT, grade_id=grade.id, is_archived=False)
    if stream is not None:
        q = q.filter_by(stream_id=stream.id)
    elif allowed_streams is not True:
        q = q.filter(User.stream_id.in_(list(allowed_streams)))
    students = q.order_by(User.first_name, User.second_name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    ws.append(["Admission Number", "Name"] + [s.name for s in sorted(subs, key=lambda x: x.name)])
    for s in students:
        ws.append([s.admission_number, s.full_name()])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 32
    how = wb.create_sheet("How to use")
    for line in ["Fill in the scores under each subject. Leave a cell empty to skip it.",
                 "Keep the Admission Number column. Learners are matched by it.",
                 "To use a different out-of mark, write it in the header, for example: Mathematics (50).",
                 "Save, then upload the file under Import marks. Nothing is saved until you confirm the preview."]:
        how.append([line])
    how.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="marks-template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
