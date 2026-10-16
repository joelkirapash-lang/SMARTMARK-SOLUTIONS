import io

from openpyxl import Workbook, load_workbook

from tests.conftest import setup_school, create_account


def _people(api, s):
    out = {}
    for first, adm in (("Amina", "1001"), ("Brian", "1002"), ("Cate", "1003")):
        j, _, _ = create_account(api, s["admin"], role="STUDENT", first_name=first, second_name="Test", admission_number=adm,
                                 grade_id=s["grade"]["id"], stream_id=s["stream"]["id"])
        out[first] = j["user"]["id"]
    return out


def _teacher(api, s, assignments):
    j, pw, tok = create_account(api, s["admin"], role="TEACHER", first_name="Tom", second_name="Ouma")
    api.post("/api/auth/change-password", tok, json={"current_password": pw, "new_password": "TeachPass1"})
    api.put(f"/api/staff/{j['user']['id']}/assignments", s["admin"], json={"assignments": assignments})
    return api.login(j["credentials"]["email"], "TeachPass1")


def _upload(api, tok, s, exam, content, filename="marks.csv", dry="1", **extra):
    data = {"exam_id": exam["id"], "grade_id": s["grade"]["id"], "stream_id": s["stream"]["id"], "dry_run": dry,
            "file": (io.BytesIO(content if isinstance(content, bytes) else content.encode()), filename), **extra}
    return api.post("/api/marks/import", tok, data=data, content_type="multipart/form-data")


def test_preview_saves_nothing_then_import_ranks_first_to_last(api, super_token):
    s = setup_school(api, super_token)
    ids = _people(api, s)
    exam = api.post("/api/exams", s["admin"], json={"name": "Term 1"}).get_json()
    csv = "Admission Number,Name,Mathematics,English\n1001,Amina,60,70\n1002,Brian,95,88\n1003,Cate,75,\n"
    r = _upload(api, s["admin"], s, exam, csv).get_json()
    assert r["dry_run"] and r["saved"] == 0
    assert r["summary"] == {"rows": 3, "matched": 3, "marks": 5, "problem_rows": 0}
    q = f"exam_id={exam['id']}&grade_id={s['grade']['id']}&subject_id={s['math']['id']}&stream_id={s['stream']['id']}"
    assert all(x["score"] is None for x in api.get(f"/api/marks/grid?{q}", s["admin"]).get_json()["rows"])
    r = _upload(api, s["admin"], s, exam, csv, dry="0").get_json()
    assert r["saved"] == 5
    ml = api.get(f"/api/marklist?exam_id={exam['id']}&grade_id={s['grade']['id']}", s["admin"]).get_json()
    assert [x["student_name"].split()[0] for x in ml["rows"]] == ["Brian", "Cate", "Amina"]
    assert [x["position"] for x in ml["rows"]] == [1, 2, 3]


def test_problems_blank_cells_names_and_custom_max(api, super_token):
    s = setup_school(api, super_token)
    ids = _people(api, s)
    exam = api.post("/api/exams", s["admin"], json={"name": "Term 1"}).get_json()
    _upload(api, s["admin"], s, exam, "Admission Number,Mathematics\n1001,80\n", dry="0")
    csv = "Name,Mathematics (50),Geography,English\nAmina Test,,10,\nBrian Test,abc,,\nCate Test,60,,\nNobody Here,5,,\n"
    r = _upload(api, s["admin"], s, exam, csv, dry="0").get_json()
    by = {c["header"]: c for c in r["subjects"]}
    assert by["Geography"]["status"] == "unknown" and by["Mathematics (50)"]["max_score"] == 50
    problems = " ".join(p for row in r["rows"] for p in row["problems"])
    assert "not a number" in problems and "outside 0 to 50" in problems and "No learner named Nobody Here" in problems
    q = f"exam_id={exam['id']}&grade_id={s['grade']['id']}&subject_id={s['math']['id']}&stream_id={s['stream']['id']}"
    rows = {x["student_name"].split()[0]: x["score"] for x in api.get(f"/api/marks/grid?{q}", s["admin"]).get_json()["rows"]}
    assert rows["Amina"] == 80          # blank cell kept the earlier mark
    assert rows["Brian"] is None and rows["Cate"] is None   # invalid and out-of-range cells were not saved


def test_teacher_limited_to_assigned_subjects_and_locked_when_published(api, super_token):
    s = setup_school(api, super_token)
    _people(api, s)
    exam = api.post("/api/exams", s["admin"], json={"name": "Term 1"}).get_json()
    tok = _teacher(api, s, [{"grade_id": s["grade"]["id"], "stream_id": s["stream"]["id"], "subject_id": s["math"]["id"]}])
    csv = "Admission Number,Mathematics,English\n1001,50,60\n"
    r = _upload(api, tok, s, exam, csv, dry="0").get_json()
    st = {c["header"]: c["status"] for c in r["subjects"]}
    assert st == {"Mathematics": "ok", "English": "blocked"} and r["saved"] == 1
    api.post(f"/api/exams/{exam['id']}/publish", s["admin"])
    assert _upload(api, tok, s, exam, csv, dry="0").status_code == 403
    assert _upload(api, s["admin"], s, exam, csv, dry="0").status_code == 200
    j, pw, stok = create_account(api, s["admin"], role="STUDENT", first_name="Zed", second_name="Kid", admission_number="9",
                                 grade_id=s["grade"]["id"])
    api.post("/api/auth/change-password", stok, json={"current_password": pw, "new_password": "StudPass1"})
    stok = api.login(j["credentials"]["email"], "StudPass1")
    assert _upload(api, stok, s, exam, csv).status_code == 403


def test_excel_file_with_numeric_admission_numbers(api, super_token):
    s = setup_school(api, super_token)
    _people(api, s)
    exam = api.post("/api/exams", s["admin"], json={"name": "Term 1"}).get_json()
    wb = Workbook()
    ws = wb.active
    ws.append(["Adm No", "Mathematics", "English"])
    ws.append([1001, 88.5, 70])
    ws.append([1002.0, 40, 41])
    buf = io.BytesIO()
    wb.save(buf)
    r = _upload(api, s["admin"], s, exam, buf.getvalue(), filename="marks.xlsx", dry="0").get_json()
    assert r["summary"]["matched"] == 2 and r["saved"] == 4


def test_template_lists_the_class_and_import_is_tenant_safe(api, super_token):
    s = setup_school(api, super_token, "Alpha School")
    other = setup_school(api, super_token, "Beta School")
    _people(api, s)
    res = api.get(f"/api/marks/import-template?grade_id={s['grade']['id']}", s["admin"])
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.data))["Marks"]
    assert [c.value for c in ws[1]] == ["Admission Number", "Name", "English", "Mathematics"]
    assert sorted(str(r[0].value) for r in ws.iter_rows(min_row=2)) == ["1001", "1002", "1003"]
    foreign_exam = api.post("/api/exams", other["admin"], json={"name": "Beta"}).get_json()
    r = _upload(api, s["admin"], s, foreign_exam, "Admission Number,Mathematics\n1001,50\n", dry="0")
    assert r.status_code == 400
